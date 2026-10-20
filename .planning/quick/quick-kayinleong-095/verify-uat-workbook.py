"""Gate for the tester UAT pack: JSON, project facts, hygiene, workbook, plan doc.

The JSON is the full catalogue (all sections). The workbook is the tester view:
one `Test Cases` sheet with the sections in meta.tester_sections.

Checks (each prints PASS or FAIL, then a total line "N PASS, M FAIL"):
  A  JSON integrity           B  project facts against the project export
  C  client-deliverable hygiene   D  workbook structure   E  plan document
A line starting NOTE is information only and is not counted.

The project export holds PIC names and phone numbers in its description field.
The loader reduces every record to five fields at once and never prints one.
The export date the facts were taken from is meta.snapshot_exported_at in the JSON;
refreshing the snapshot means editing that value on purpose.

Usage:
  python3 verify-uat-workbook.py [--cases P] [--workbook P] [--plan P] [--export P]
  python3 verify-uat-workbook.py --self-test    (proves the gate bites)
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError:
    sys.exit("openpyxl is not importable. Stop and report; do not install anything.")

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
DEF_CASES = HERE / "uat-cases.json"
DEF_WORKBOOK = ROOT / "docs" / "UAT-Output-Test-Cases.xlsx"
DEF_PLAN = ROOT / "docs" / "UAT-Output-Test-Plan.md"
DEF_EXPORT = ROOT / "projects.export.json"

LANGS = ["EN", "BM", "中文"]
TYPES = ["Golden", "Rubric", "Adversarial"]
EXPORT_KEYS = ("name", "status", "priceValue", "foreignEligible", "bumiQuota")

SHEETS = ["Test Cases"]
TC_HEADERS = ["ID", "Area", "Test Case", "Priority", "Preconditions", "Steps", "Expected Result", "Result",
              "Actual / Comments", "Known Status", "Lang", "Type", "Runs", "Run by", "Double-rate"]
TC_WIDTHS = [7, 16, 26, 9, 24, 40, 44, 12, 30, 34, 8, 13, 7, 15, 11]
RESULT_LIST = '"Pass,Fail,Blocked,Not Run"'

PLAN_HEADINGS = ["## 1. Purpose and scope", "## 2. What is being tested", "## 3. Roles",
                 "## 4. Before we start: readiness checklist", "## 5. Schedule",
                 "## 6. How to run a case on your phone", "## 7. How to judge an answer",
                 "## 8. How serious is it? (severity)", "## 9. What to record", "## 10. Exit criteria",
                 "## 11. Known issues", "## 12. Open questions for Derek", "## 13. Appendix: case counts"]
# Known-status key -> words that must appear in plan section 11 when a tester case carries the key.
KEY_PROBES = {
    "KS-RM": ["RM"], "KS-TOAST": ["notified"], "KS-TABLE": ["brochures"], "KS-SEGMENT": ["Highlight"],
    "KS-UI-EN": ["English"], "KS-CUTOFF": ["cut off"], "KS-BADGE": ["Cold prospect"], "KS-TOKENS": ["<LEAD_ID"],
    "KS-CJK": ["Chinese place names"], "KS-BUMI": ["bumi"], "KS-SOLDOUT": ["sold out"],
    "KS-PROPERTY1": ["Property 1"], "KS-MIXED-DATA": ["Mixed write-ups"], "KS-ADS": ["Paid ads"],
    "KS-PLACEHOLDER": ["example SOP"],
}
EXTRA_PROBES = [["previous conversation", "across conversations"]]
# Issues that only the removed sections touch: they belong in the "Not tested this round" note.
NOT_TESTED_PROBES = [["phone numbers"], ["Reply in"], ["journey"]]
REMOVED_TAB_RE = re.compile(r"Run Log|Sign-off|Rubric tab|Instructions tab|Cross-cutting|\bX\d+\b|(?<!Type )\bE\d+\b|E1.E50")  # "Type E1" is a layout name, not a slot

ALLOW = {"012-000 0000", "0120000000", "+60 12-000 0000", "900101-14-0000", "900101140000", "siti@example.com"}
ALLOW_NAMES = {"Encik Ali", "Cik Tan", "Ms Tan", "Mr Lim", "Madam Wong", "Puan Wong", "Mr Tanaka", "Mr Lee",
               "Mr Chen"}
PHONE_RE = re.compile(r"(?<![\d,.])(?:\+?6?0[\s-]?)?1\d[\s-]?\d{3,4}[\s-]?\d{4}(?![\d,])")
IC_RE = re.compile(r"(?<!\d)\d{6}-?\d{2}-?\d{4}(?!\d)")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
LONG_RE = re.compile(r"(?<![\d,.])\d{9,}(?![\d,.])")
NAME_RE = re.compile(r"\b(?:Encik|Cik|Puan|Ms|Mr|Mrs|Madam) [A-Z][a-z]+")
SECRET_RE = re.compile(r"sk-ant-[A-Za-z0-9_-]{10,}|AIza[0-9A-Za-z_-]{30,}|-----BEGIN [A-Z ]*PRIVATE KEY-----")
FORBIDDEN_RE = re.compile(r"\b(?:src|app|scripts|evals|tests)/[\w\[.-]|\.(?:tsx?|py|json|yaml)\b|RESEARCH-|"
                          r"quick-kayinleong|projects\.export|gitignore")
ID_RE = re.compile(r"^[FRCX][1-9]\d*$")

FAILURES = []
PASSES = []


def check(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    (PASSES if ok else FAILURES).append(msg)
    return ok


def short(items, n=4):
    items = list(items)
    s = "; ".join(str(i) for i in items[:n])
    return s + (" (+{} more)".format(len(items) - n) if len(items) > n else "")


def ceil_pct(n, pct):
    return -(-n * pct // 100)


def iter_strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            for s in iter_strings(k):
                yield s
            for s in iter_strings(v):
                yield s
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            for s in iter_strings(v):
                yield s


def compose_expected(case):
    e = case["expected"]
    return "MUST: {}\nMUST NOT: {}\nALSO ACCEPTABLE: {}".format(e["must"], e["must_not"], e["also_acceptable"])


def hygiene_hits(text):
    hits = []
    for rx in (PHONE_RE, IC_RE, EMAIL_RE, LONG_RE):
        hits.extend(m.group(0) for m in rx.finditer(text))
    return [h for h in hits if h not in ALLOW]


def name_hits(text):
    return [m.group(0) for m in NAME_RE.finditer(text) if m.group(0) not in ALLOW_NAMES]


def load_export(path):
    """Reduce every record to the five allowed fields at once; never print a record."""
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    exported_at = raw.get("exportedAt")
    reduced = []
    for rec in raw.get("records", []):
        reduced.append({k: rec.get(k) for k in EXPORT_KEYS})
    del raw
    for rec in reduced:
        assert set(rec.keys()) == set(EXPORT_KEYS)
    return exported_at, reduced


def cell_text(v):
    return None if v is None else str(v)


def note(msg):
    print("NOTE " + msg)


def tester_letters(data):
    return list(data["meta"].get("tester_sections", []))


def tester_sections(data):
    letters = tester_letters(data)
    return [s for s in data["sections"] if s["letter"] in letters]


def tester_cases(data):
    letters = tester_letters(data)
    return [c for c in data["cases"] if c["section"] in letters]


def count_rows(data):
    """Rows of the plan's appendix table, as lists of strings, from the tester cases."""
    cases = tester_cases(data)
    rows = [["Section", "Language", "Golden", "Rubric", "Adversarial", "Total"]]
    for s in tester_sections(data):
        for lang in LANGS:
            sub = [c for c in cases if c["section"] == s["letter"] and c["lang"] == lang]
            cnt = Counter(c["type"] for c in sub)
            rows.append([s["pillar"], lang] + [str(cnt.get(t, 0)) for t in TYPES] + [str(len(sub))])
    cnt = Counter(c["type"] for c in cases)
    rows.append(["All sections", "All"] + [str(cnt.get(t, 0)) for t in TYPES] + [str(len(cases))])
    return rows


def needed_rows(data):
    """Rows of the exact-numbers table in plan section 10 (EC3, EC4, EC5, EC6)."""
    cases = tester_cases(data)
    rows = [["Gate", "Cases", "Needed to pass"]]
    for s in tester_sections(data):
        n = sum(1 for c in cases if c["section"] == s["letter"])
        rows.append(["EC3 {}".format(s["pillar"]), str(n), str(ceil_pct(n, 85))])
    for lang in LANGS:
        n = sum(1 for c in cases if c["lang"] == lang)
        rows.append(["EC4 {} (all pillars)".format(lang), str(n), str(ceil_pct(n, 80))])
    for s in tester_sections(data):
        for lang in LANGS:
            n = sum(1 for c in cases if c["section"] == s["letter"] and c["lang"] == lang)
            if n >= 10:
                rows.append(["EC4 {}, {} only (cell of 10 or more)".format(s["pillar"], lang), str(n),
                             str(ceil_pct(n, 80))])
    adv = sum(1 for c in cases if c["type"] == "Adversarial")
    rows.append(["EC5 Adversarial cases", str(adv), str(ceil_pct(adv, 90))])
    dr = sum(1 for c in cases if c["double_rate"])
    rows.append(["EC6 Double-rated cases (raters agree on)", str(dr), str(ceil_pct(dr, 80))])
    return rows


def md_tables(text):
    """Every markdown table in text, each as a list of rows of stripped cells (rule rows dropped)."""
    tables, cur = [], []
    for line in text.splitlines():
        if line.startswith("|"):
            cells = [x.strip() for x in line.strip().strip("|").split("|")]
            if not all(re.match(r"^:?-+:?$", x) for x in cells):
                cur.append(cells)
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    return tables


def section_of(plan_text, start_heading, end_heading=None):
    a = plan_text.find(start_heading)
    if a < 0:
        return ""
    b = plan_text.find(end_heading, a) if end_heading else -1
    return plan_text[a:b] if b >= 0 else plan_text[a:]


# ------------------------------------------------------------------ A: JSON
def run_a(data):
    cases = data["cases"]
    n = len(cases)
    check(n == 77 and 60 <= n <= 80, "A1 scripted case count is 77 and within 60-80 (got {})".format(n))

    ids = [c["id"] for c in cases]
    dup = [i for i, k in Counter(ids).items() if k > 1]
    check(not dup, "A2 IDs are unique (duplicates: {})".format(short(dup) or "none"))
    bad_fmt = [c["id"] for c in cases if not ID_RE.match(c["id"])
               or c["id"][0] != c["section"]]
    check(not bad_fmt, "A2 IDs match the pattern and their section letter ({})".format(short(bad_fmt) or "ok"))
    seq_bad = []
    for letter in [s["letter"] for s in data["sections"]]:
        nums = []
        for c in cases:
            if c["section"] == letter and ID_RE.match(c["id"]):
                nums.append(int(c["id"][1:]))
        if nums != list(range(1, len(nums) + 1)):
            seq_bad.append(letter)
    check(not seq_bad, "A2 IDs are sequential from 1 within each section (broken: {})".format(
        short(seq_bad) or "none"))
    order_ok = [c["section"] for c in cases] == sorted(
        [c["section"] for c in cases], key=lambda x: [s["letter"] for s in data["sections"]].index(x))
    check(order_ok, "A2 cases are listed section by section in workbook order")

    blk_bad = []
    for s in data["sections"]:
        langs = [c["lang"] for c in cases if c["section"] == s["letter"]]
        collapsed = [langs[i] for i in range(len(langs)) if i == 0 or langs[i] != langs[i - 1]]
        if collapsed != LANGS:
            blk_bad.append(s["letter"])
    check(not blk_bad, "A3 each section has contiguous EN, BM, 中文 blocks in that order (broken: {})".format(
        short(blk_bad) or "none"))

    cnt = Counter(c["type"] for c in cases)
    shares = {t: cnt.get(t, 0) * 100.0 / n for t in TYPES}
    check(40 <= shares["Golden"] <= 50 and 15 <= shares["Rubric"] <= 25 and 25 <= shares["Adversarial"] <= 40,
          "A4 type shares within bands (Golden {:.0f}%, Rubric {:.0f}%, Adversarial {:.0f}%)".format(
              shares["Golden"], shares["Rubric"], shares["Adversarial"]))
    runs_bad = [c["id"] for c in cases if c["runs"] != (3 if c["type"] == "Adversarial" else 2)]
    check(not runs_bad, "A5 runs is 3 for Adversarial and 2 otherwise ({})".format(short(runs_bad) or "ok"))

    # A6: required fields, allowed sets
    a6 = []
    required = ["id", "section", "lang", "type", "priority", "area", "title", "preconditions", "steps", "expected",
                "known_status_keys", "runs", "run_by", "double_rate", "repoint_when_real_sops", "projects",
                "source"]
    for c in cases:
        cid = c.get("id", "?")
        for k in required:
            if k not in c:
                a6.append("{} missing {}".format(cid, k))
        for k in ["id", "section", "lang", "type", "priority", "area", "title", "preconditions", "steps", "run_by",
                  "source"]:
            if not (isinstance(c.get(k), str) and c.get(k).strip()):
                a6.append("{} empty {}".format(cid, k))
        exp = c.get("expected") or {}
        for k in ["must", "must_not", "also_acceptable"]:
            if not (isinstance(exp.get(k), str) and exp.get(k).strip()):
                a6.append("{} empty expected.{}".format(cid, k))
        if not c.get("known_status_keys"):
            a6.append("{} no known_status_keys".format(cid))
        if c.get("lang") not in LANGS:
            a6.append("{} bad lang".format(cid))
        if c.get("type") not in TYPES:
            a6.append("{} bad type".format(cid))
        if c.get("priority") not in ("P0", "P1", "P2"):
            a6.append("{} bad priority".format(cid))
        if not isinstance(c.get("double_rate"), bool):
            a6.append("{} double_rate not boolean".format(cid))
        want = None
        if c.get("lang") == "BM":
            want = "BM speaker"
        elif c.get("lang") == "中文":
            want = "中文 speaker"
        elif c.get("section") in ("F", "X"):
            want = "New agent"
        elif c.get("section") in ("R", "C"):
            want = "Senior coach"
        if c.get("run_by") != want:
            a6.append("{} run_by {!r} should be {!r}".format(cid, c.get("run_by"), want))
    check(not a6, "A6 every field present and non-empty with allowed values ({})".format(short(a6) or "ok"))

    a7 = [c["id"] for c in cases if not re.search(r'"[^"\n]{2,}"', c["steps"])]
    check(not a7, "A7 steps contain a double-quoted prompt ({})".format(short(a7) or "ok"))

    # A8
    ph = data["known_status_phrases"]
    a8 = []
    for c in cases:
        keys = c["known_status_keys"]
        steps = c["steps"]
        if re.search(r"RM\d", steps) and "KS-RM" not in keys:
            a8.append("{} needs KS-RM".format(c["id"]))
        if any(p in steps for p in ("012-000 0000", "0120000000", "900101140000")) and "KS-PHONE" not in keys:
            a8.append("{} needs KS-PHONE".format(c["id"]))
        if c["section"] == "R" and c["lang"] in ("BM", "中文") and "KS-BADGE" not in keys:
            a8.append("{} needs KS-BADGE".format(c["id"]))
        if c["section"] == "C" and c["type"] == "Golden" and not (
                "KS-PLACEHOLDER" in keys and c["repoint_when_real_sops"] is True):
            a8.append("{} needs KS-PLACEHOLDER and repoint".format(c["id"]))
        if c["area"] == "Coach / KB miss" and "KS-TOAST" not in keys:
            a8.append("{} needs KS-TOAST".format(c["id"]))
        for k in keys:
            if k not in ph:
                a8.append("{} unknown key {}".format(c["id"], k))
        if "KS-OK" in keys and len(keys) > 1:
            a8.append("{} KS-OK combined".format(c["id"]))
    check(not a8, "A8 known-status rules hold ({})".format(short(a8) or "ok"))

    dr = sum(1 for c in cases if c["double_rate"])
    dr_t = sum(1 for c in tester_cases(data) if c["double_rate"])
    check(12 <= dr <= 18 and 12 <= dr_t <= 18,
          "A9 double-rated cases within 12-18 (catalogue {}, tester view {})".format(dr, dr_t))

    a10 = []
    if [r["code"] for r in data["rubric"]] != ["H1", "H2", "H3", "H4", "H5", "H6", "H7", "U"]:
        a10.append("rubric codes")
    for r in data["rubric"]:
        for k in ("question", "pass", "fail"):
            if not str(r.get(k, "")).strip():
                a10.append("{} empty {}".format(r.get("code"), k))
    if [s["level"] for s in data["severity"]] != ["S1", "S2", "S3", "S4"]:
        a10.append("severity levels")
    if len(data["root_causes"]) != 5:
        a10.append("root causes count")
    if [e["id"] for e in data["exit_criteria"]] != ["EC{}".format(i) for i in range(1, 8)]:
        a10.append("exit criteria ids")
    check(not a10, "A10 rubric, severity, root causes and exit criteria are complete ({})".format(
        short(a10) or "ok"))

    letters = tester_letters(data)
    known = [s["letter"] for s in data["sections"]]
    a11 = []
    if letters != ["F", "R", "C"]:
        a11.append("tester_sections is {}".format(letters))
    if any(x not in known for x in letters) or letters != [x for x in known if x in letters]:
        a11.append("tester_sections is not a subset of the sections, in workbook order")
    if len(tester_cases(data)) != 63:
        a11.append("{} tester cases".format(len(tester_cases(data))))
    check(not a11, "A11 meta.tester_sections is F, R, C and covers 63 cases ({})".format(short(a11) or "ok"))


# ------------------------------------------------------------------ B: project facts
def run_b(data, export_path):
    try:
        exported_at, recs = load_export(export_path)
    except Exception as exc:  # message only, never content
        check(False, "B1 export loads ({})".format(type(exc).__name__))
        return
    snap = data["meta"].get("snapshot_exported_at")
    check(bool(snap) and exported_at == snap and len(recs) > 0,
          "B1 export loads ({} records reduced to 5 fields) and exportedAt matches meta.snapshot_exported_at "
          "(refresh the snapshot by editing that value)".format(len(recs)))
    by_name = {}
    for r in recs:
        by_name.setdefault(r["name"], []).append(r)

    declared = []
    for c in data["cases"]:
        for p in c["projects"]:
            declared.append((c["id"], p))
    b2 = [(cid, p["name"]) for cid, p in declared if len(by_name.get(p["name"], [])) != 1]
    check(not b2, "B2 every declared project name exists exactly once in the export ({})".format(
        short(["{}: {}".format(a, b) for a, b in b2]) or "{} declared".format(len(declared))))

    b3, b4, b5 = [], [], []
    for cid, p in declared:
        rows = by_name.get(p["name"], [])
        if len(rows) != 1:
            continue
        r = rows[0]
        e = p["expect"]
        pv = r["priceValue"] or 0
        if (e["status"] != r["status"] or abs(float(e["priceValue"]) - float(pv)) > 1
                or bool(e["foreignEligible"]) != bool(r["foreignEligible"])
                or bool(e["bumiQuota"]) != bool(r["bumiQuota"])):
            b3.append("{}: {}".format(cid, p["name"]))
        if p["clause"] == "must" and r["status"] != "active":
            b4.append("{}: {} is {} in the export".format(cid, p["name"], r["status"]))
        case = next(c for c in data["cases"] if c["id"] == cid)
        if p["name"] not in case["expected"].get(p["clause"], ""):
            b5.append("{}: {}".format(cid, p["name"]))
    check(not b3, "B3 every expect value equals the export ({})".format(short(b3) or "ok"))
    check(not b4, "B4 projects under MUST are active in the export ({})".format(short(b4) or "ok"))
    check(not b5, "B5 every declared name appears verbatim in its clause text ({})".format(short(b5) or "ok"))

    names = sorted((n for n in by_name if n), key=len, reverse=True)
    b6 = []
    for c in data["cases"]:
        decl = {(p["name"], p["clause"]) for p in c["projects"]}
        for clause in ("must", "must_not", "also_acceptable"):
            text = c["expected"].get(clause, "")
            masked = text
            for n in names:
                if n in masked:
                    if (n, clause) not in decl:
                        b6.append("{}: {} in {}".format(c["id"], n, clause))
                    masked = masked.replace(n, "\0" * len(n))
    check(not b6, "B6 every export project name found in a clause is declared for that clause ({})".format(
        short(b6) or "ok"))


# ------------------------------------------------------------------ workbook cells
def workbook_cells(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    yield ws.title, c


# ------------------------------------------------------------------ C: hygiene
def run_c(data, wb, plan_text):
    wb_strings = [(ws_name + "!" + c.coordinate, str(c.value)) for ws_name, c in workbook_cells(wb)]
    json_strings = [("json", s) for s in iter_strings(data)]
    plan_strings = [("plan", plan_text)]
    everything = json_strings + wb_strings + plan_strings

    c1 = []
    for where, s in everything:
        for h in hygiene_hits(s):
            c1.append("{} contains {!r}".format(where, h))
    check(not c1, "C1 every phone, IC, email and long-digit match is a placeholder ({})".format(
        short(c1) or "none outside the allowlist"))
    c2 = []
    for where, s in everything:
        for h in name_hits(s):
            c2.append("{} has {!r}".format(where, h))
    check(not c2, "C2 every honorific name is a synthetic name ({})".format(short(c2) or "ok"))

    c3 = []
    for where, s in wb_strings + plan_strings:
        m = FORBIDDEN_RE.search(s)
        if m:
            c3.append("{} has {!r}".format(where, m.group(0)))
    src_values = [c["source"] for c in data["cases"]]
    for where, s in wb_strings:
        for sv in src_values:
            if sv in s:
                c3.append("{} contains a source label".format(where))
                break
    check(not c3, "C3 no internal reference, path or source label in the plan doc or any workbook cell ({})".format(
        short(c3) or "ok"))

    c4 = []
    for where, s in everything:
        if SECRET_RE.search(s):
            c4.append(where)
    check(not c4, "C4 no secrets in any deliverable ({})".format(short(c4) or "ok"))


# ------------------------------------------------------------------ D: workbook structure
def dv_cells(ws):
    """Map each data-validation formula to the set of cells it covers."""
    out = {}
    for dv in ws.data_validations.dataValidation:
        cells = out.setdefault(dv.formula1, set())
        for rng in str(dv.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rng)
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    cells.add((rr, cc))
    return out


def fill_of(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type == "solid" else None


def font_color(cell):
    col = cell.font.color
    return col.rgb[-6:] if col is not None and isinstance(col.rgb, str) else None


def analyse_test_cases(ws):
    """Derive the ID rows from the sheet itself (independent of the builder)."""
    rows = {}
    for r in range(2, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and re.match(r"^[A-Z][1-9]\d*$", a):
            rows[a] = r
    return rows


def run_d(data, wb):
    cases = tester_cases(data)
    sections = tester_sections(data)
    phrases = data["known_status_phrases"]

    check(wb.sheetnames == SHEETS, "D1 the workbook has one sheet, {} (got {})".format(SHEETS, wb.sheetnames))
    if wb.sheetnames != SHEETS:
        return
    tc = wb["Test Cases"]
    size = (tc.max_row, tc.max_column)  # read first: ws.cell() on an empty cell grows the sheet

    heads = [tc.cell(row=1, column=c).value for c in range(1, 16)]
    check(heads == TC_HEADERS and tc.cell(row=1, column=16).value is None,
          "D2 Test Cases headers A-O are exact (got {})".format(heads))

    id_rows = analyse_test_cases(tc)
    band_texts = [s["band"] for s in sections]
    band_rows = {}
    for r in range(2, tc.max_row + 1):
        a = tc.cell(row=r, column=1).value
        if a in band_texts:
            band_rows[a] = r
    merged = {str(m) for m in tc.merged_cells.ranges}
    d3 = []
    for c in range(1, 16):
        h = tc.cell(row=1, column=c)
        if fill_of(h) != "1F2937" or font_color(h) != "FFFFFF" or not h.font.b:
            d3.append("header col {}".format(c))
    if len(band_rows) != len(sections):
        d3.append("expected {} bands, found {}".format(len(sections), len(band_rows)))
    if merged != {"A{r}:O{r}".format(r=r) for r in band_rows.values()}:
        d3.append("merged ranges are {}".format(sorted(merged)))
    for text, r in band_rows.items():
        cell = tc.cell(row=r, column=1)
        if fill_of(cell) != "D9F99D" or not cell.font.b:
            d3.append("band row {} style".format(r))
    for ident, r in id_rows.items():
        for col in (8, 9):
            if fill_of(tc.cell(row=r, column=col)) != "FFF9C4":
                d3.append("{} col {} not yellow".format(ident, col))
        if font_color(tc.cell(row=r, column=10)) != "6B7280":
            d3.append("{} Known Status colour".format(ident))
    if tc.freeze_panes != "A2":
        d3.append("freeze panes {}".format(tc.freeze_panes))
    for i, w in enumerate(TC_WIDTHS, start=1):
        got = tc.column_dimensions[get_column_letter(i)].width
        if got is None or abs(got - w) > 0.01:
            d3.append("width {} is {}".format(get_column_letter(i), got))
    fonts = {c.font.name for _, c in workbook_cells(wb)}
    if fonts != {"Arial"}:
        d3.append("fonts {}".format(sorted(fonts)))
    if size != (len(sections) + len(cases) + 1, 15):
        d3.append("sheet is {} rows by {} columns".format(*size))
    check(not d3, "D3 header and band fills, merges, yellow cells, grey J, freeze A2, widths, Arial, size ({})".format(
        short(d3) or "ok"))

    dvs = dv_cells(tc)
    covered = dvs.get(RESULT_LIST, set())
    want = {(r, 8) for r in id_rows.values()}
    d4 = []
    if covered != want:
        d4.append("covered {} cells, wanted {} (missing {}, extra {})".format(
            len(covered), len(want), len(want - covered), len(covered - want)))
    if set(dvs) - {RESULT_LIST}:
        d4.append("other validations present")
    if any((r, 8) in covered for r in band_rows.values()):
        d4.append("a band row is covered")
    notrun = [i for i, r in id_rows.items() if tc.cell(row=r, column=8).value != "Not Run"]
    if notrun:
        d4.append("H not 'Not Run': {}".format(short(notrun)))
    if len(id_rows) != len(cases):
        d4.append("{} ID rows, wanted {}".format(len(id_rows), len(cases)))
    check(not d4, "D4 Result validation is exactly {} on every case row, no band row, all Not Run ({})".format(
        RESULT_LIST, short(d4) or "ok"))

    seq = []
    for s in sections:
        seq.append(s["band"])
        seq.extend(c["id"] for c in cases if c["section"] == s["letter"])
    got_seq = [tc.cell(row=r, column=1).value for r in range(2, tc.max_row + 1)]
    d5 = []
    if got_seq != seq:
        d5.append("row order differs from JSON (first diff at index {})".format(
            next((i for i, (x, y) in enumerate(zip(got_seq, seq)) if x != y), min(len(got_seq), len(seq)))))
    for c in cases:
        r = id_rows.get(c["id"])
        if r is None:
            d5.append("{} missing".format(c["id"]))
            continue
        exp = [c["id"], c["area"], c["title"], c["priority"], c["preconditions"], c["steps"], compose_expected(c),
               "Not Run", None, "\n".join(phrases[k] for k in c["known_status_keys"]), c["lang"], c["type"],
               c["runs"], c["run_by"], "Yes" if c["double_rate"] else None]
        got = [tc.cell(row=r, column=k).value for k in range(1, 16)]
        for k, (g, e) in enumerate(zip(got, exp), start=1):
            if g != e:
                d5.append("{} {} differs".format(c["id"], get_column_letter(k)))
    check(not d5, "D5 band and case rows in JSON order, every case row equals the JSON-composed values ({})".format(
        short(d5) or "{} rows ok".format(len(cases))))

    d6 = []
    for row in tc.iter_rows():
        for c in row:
            if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                d6.append(c.coordinate)
    left_out = [i for i in id_rows if i[0] not in tester_letters(data)]
    if left_out:
        d6.append("rows from sections not in the tester view: {}".format(short(left_out)))
    check(not d6, "D6 no formulas and no cross-cutting or exploratory rows ({})".format(short(d6) or "ok"))

    run_log = sum(1 for _, c in workbook_cells(wb) if "Run Log" in str(c.value))
    if run_log:
        note("{} workbook cells (the Steps of each case) still say 'copy the full answer into the Run Log'; "
             "that tab is gone. The case text comes from the JSON and is not changed here.".format(run_log))


# ------------------------------------------------------------------ E: plan doc
def run_e(data, plan_text):
    cases = tester_cases(data)
    pos = [plan_text.find(h) for h in PLAN_HEADINGS]
    check(all(p >= 0 for p in pos) and pos == sorted(pos),
          "E1 the 13 plan headings are present in order (positions {})".format(pos))

    # E2: appendix
    appendix = section_of(plan_text, "## 13. Appendix: case counts")
    tables = md_tables(appendix)
    table = tables[0] if tables else []
    want = count_rows(data)
    e2 = []
    if table != want:
        bad = [i for i in range(max(len(table), len(want)))
               if i >= len(table) or i >= len(want) or table[i] != want[i]]
        e2.append("table rows differ at {}".format(bad[:6]))
    runs = sum(c["runs"] for c in cases)
    dr = sum(1 for c in cases if c["double_rate"])
    lines = [
        "Planned runs: {} (2 per Golden or Rubric case, 3 per Adversarial case).".format(runs),
        "Double-rated cases: {}.".format(dr),
    ]
    for ln in lines:
        if ln not in appendix:
            e2.append("summary line missing: {}".format(ln[:30]))
    if re.search(r"[Ee]xploratory", appendix):
        e2.append("appendix still mentions exploratory slots")
    check(not e2, "E2 appendix table (Finder, Reply, Coach only) and the two summary lines equal the JSON counts "
          "({})".format(short(e2) or "{} table rows ok".format(len(want))))

    # E3: section 10
    s10 = section_of(plan_text, "## 10. Exit criteria", "## 11. Known issues")
    e3 = []
    if "Proposed — Derek to confirm" not in s10:
        e3.append("Proposed label missing")
    if "Cross-cutting" in s10:
        e3.append("Cross-cutting still in section 10")
    got = next((t for t in md_tables(s10) if t and t[0] == ["Gate", "Cases", "Needed to pass"]), None)
    exp = needed_rows(data)
    if got != exp:
        bad = [i for i in range(max(len(got or []), len(exp)))
               if got is None or i >= len(got) or i >= len(exp) or got[i] != exp[i]]
        e3.append("exact-numbers table differs at {}".format(bad[:6]))
    check(not e3, "E3 section 10 keeps 'Proposed — Derek to confirm' and its exact-numbers table equals the JSON "
          "counts ({})".format(short(e3) or "{} rows ok".format(len(exp) - 1)))

    # E4: known issues
    s11 = section_of(plan_text, "## 11. Known issues", "## 12.")
    low = s11.lower()
    present = {k for c in cases for k in c["known_status_keys"]}
    probes = [(k, KEY_PROBES[k]) for k in sorted(present) if k in KEY_PROBES]
    probes += [("extra", p) for p in EXTRA_PROBES]
    miss = [k + ":" + " / ".join(p) for k, p in probes if not any(w.lower() in low for w in p)]
    unknown = sorted(k for k in present if k not in KEY_PROBES and k != "KS-OK")
    if unknown:
        miss.append("no probe defined for {}".format(unknown))
    m = re.search(r"\*\*Not tested this round\.?\*\*", s11)
    tail = low[m.start():] if m else ""
    if not m:
        miss.append("'Not tested this round' note missing")
    else:
        for p in NOT_TESTED_PROBES:
            if not any(w.lower() in tail for w in p):
                miss.append("not-tested note lacks " + " / ".join(p))
        head = low[:m.start()]
        for p in NOT_TESTED_PROBES:
            if any(w.lower() in head for w in p):
                miss.append("{} is still listed as a tested known issue".format(" / ".join(p)))
    check(not miss, "E4 section 11 covers every known issue on the Finder, Reply and Coach cases, and moves the rest "
          "to 'Not tested this round' (problems: {})".format(short(miss) or "none"))

    # E5: plain English
    m = FORBIDDEN_RE.search(plan_text)
    check(m is None, "E5 plan doc is plain English, no paths or internal labels ({})".format(
        m.group(0) if m else "ok"))

    # E6: removed tabs and sections
    m = REMOVED_TAB_RE.search(plan_text)
    check(m is None, "E6 plan doc names no removed tab, cross-cutting case or exploratory slot ({})".format(
        repr(m.group(0)) if m else "ok"))

    # E7: rubric anchors in section 7
    s7 = section_of(plan_text, "## 7. How to judge an answer", "## 8.")
    rows = [r for t in md_tables(s7) for r in t]
    e7 = []
    for item in data["rubric"]:
        want_row = [item["code"], item["name"], item["question"], item["pass"], item["fail"], item["applies_to"]]
        if want_row not in rows:
            e7.append(item["code"])
    hdr = ["Code", "Name", "Ask yourself", "PASS looks like", "FAIL looks like", "Applies to"]
    if hdr not in rows:
        e7.append("header")
    for rule in data["pass_rules"]:
        label = rule.partition(":")[0]
        if label not in s7:
            e7.append("pass rule " + label)
    check(not e7, "E7 section 7 holds the H1-H7 and U anchors, PASS and FAIL for each, equal to the JSON ({})".format(
        short(e7) or "8 rows ok"))

    # E8: schedule and recording
    s5 = section_of(plan_text, "## 5. Schedule", "## 6.")
    s6 = section_of(plan_text, "## 6. How to run", "## 7.")
    s9 = section_of(plan_text, "## 9. What to record", "## 10.")
    e8 = []
    if "{} runs".format(runs) not in s5 or "{} cases".format(len(cases)) not in s5:
        e8.append("section 5 load line lacks {} runs / {} cases".format(runs, len(cases)))
    for name, sec in (("6", s6), ("9", s9)):
        for w in ("Result", "Actual / Comments", "screenshot", "UAT coordinator"):
            if w.lower() not in sec.lower():
                e8.append("section {} lacks {!r}".format(name, w))
    s3 = section_of(plan_text, "## 3. Roles", "## 4.")
    if re.search(r"Exploratory|E1", s3):
        e8.append("section 3 still has exploratory rows")
    check(not e8, "E8 sections 3, 5, 6 and 9 match the tester view ({})".format(short(e8) or "ok"))


# ------------------------------------------------------------------ main
def verify(cases_path, workbook_path, plan_path, export_path):
    del FAILURES[:]
    del PASSES[:]
    data = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    plan_text = Path(plan_path).read_text(encoding="utf-8")
    run_a(data)
    run_b(data, export_path)
    wb = load_workbook(str(workbook_path))
    run_c(data, wb, plan_text)
    run_d(data, wb)
    run_e(data, plan_text)
    print("\n{} PASS, {} FAIL".format(len(PASSES), len(FAILURES)))
    return 1 if FAILURES else 0


def self_test(args):
    script = str(Path(__file__).resolve())
    ok_all = True
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def stage(name):
            d = tmp / name
            d.mkdir()
            cj, wx, pl = d / "uat-cases.json", d / "wb.xlsx", d / "plan.md"
            shutil.copy(args.cases, cj)
            shutil.copy(args.workbook, wx)
            shutil.copy(args.plan, pl)
            return cj, wx, pl

        def run(cj, wx, pl):
            p = subprocess.run([sys.executable, script, "--cases", str(cj), "--workbook", str(wx), "--plan",
                                str(pl), "--export", str(args.export)], capture_output=True, text=True)
            return p.returncode, p.stdout + p.stderr

        cj, wx, pl = stage("control")
        rc, out = run(cj, wx, pl)
        print("{} control copy passes (exit {})".format("OK  " if rc == 0 else "FAIL", rc))
        ok_all &= rc == 0
        if rc != 0:
            print("\n".join(l for l in out.splitlines() if l.startswith("FAIL")))

        def edit_json(name, mutate):
            cj, wx, pl = stage(name)
            d = json.loads(cj.read_text(encoding="utf-8"))
            mutate(d)
            cj.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
            return cj, wx, pl

        def case(d, cid):
            return next(c for c in d["cases"] if c["id"] == cid)

        def m1(d):
            case(d, "R2")["id"] = "R1"

        def m2(d):
            c = case(d, "F5")
            c["expected"]["must"] += " Papyrus North Kiara"
            c["projects"].append({"name": "Papyrus North Kiara", "clause": "must",
                                  "expect": {"status": "hidden", "priceValue": 760000, "foreignEligible": False,
                                             "bumiQuota": False}})

        def m3(d):
            c = case(d, "R2")
            c["steps"] += " Call 012-345 6789."

        def m4(d):
            case(d, "C3")["expected"]["must_not"] = ""

        defects = [
            ("duplicate ID", edit_json("d1", m1), "FAIL A2 IDs are unique"),
            ("hidden project under MUST", edit_json("d2", m2), "FAIL B4"),
            ("non-placeholder phone", edit_json("d3", m3), "FAIL C1"),
            ("empty MUST NOT", edit_json("d4", m4), "FAIL A6"),
        ]
        cj, wx, pl = stage("d5")
        wb = load_workbook(str(wx))
        wb["Test Cases"]["J1"] = "Status"
        wb.save(str(wx))
        defects.append(("changed header", (cj, wx, pl), "FAIL D2"))

        cj, wx, pl = stage("d6")
        d = json.loads(cj.read_text(encoding="utf-8"))
        rows = count_rows(d)
        total = rows[-1]
        needle = "| " + " | ".join(total) + " |"
        bumped = total[:-1] + [str(int(total[-1]) + 1)]
        text = pl.read_text(encoding="utf-8")
        if needle in text:
            pl.write_text(text.replace(needle, "| " + " | ".join(bumped) + " |"), encoding="utf-8")
        else:
            print("FAIL appendix total row not found in the plan copy")
            ok_all = False
        defects.append(("changed appendix total", (cj, wx, pl), "FAIL E2"))

        cj, wx, pl = stage("d7")
        wb = load_workbook(str(wx))
        wb.create_sheet("Run Log")
        wb.save(str(wx))
        defects.append(("extra sheet in the workbook", (cj, wx, pl), "FAIL D1"))

        cj, wx, pl = stage("d8")
        pl.write_text(pl.read_text(encoding="utf-8") + "\nSee the Run Log tab for each answer.\n", encoding="utf-8")
        defects.append(("plan names a removed tab", (cj, wx, pl), "FAIL E6"))

        cj, wx, pl = stage("d9")
        wb = load_workbook(str(wx))
        ws = wb["Test Cases"]
        ws.freeze_panes = "A19"
        wb.save(str(wx))
        defects.append(("freeze panes moved", (cj, wx, pl), "FAIL D3"))

        for name, paths, expect in defects:
            rc, out = run(*paths)
            bites = rc != 0 and expect in out
            print("{} {}{}".format("BITES" if bites else "MISSED", name,
                                   "" if bites else " (exit {}, expected {!r})".format(rc, expect)))
            ok_all &= bites
    print("\nself-test {}".format("PASSED" if ok_all else "FAILED"))
    return 0 if ok_all else 1


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", default=str(DEF_CASES))
    ap.add_argument("--workbook", default=str(DEF_WORKBOOK))
    ap.add_argument("--plan", default=str(DEF_PLAN))
    ap.add_argument("--export", default=str(DEF_EXPORT), help="project export JSON (default: projects.export.json)")
    ap.add_argument("--self-test", action="store_true")
    args = ap.parse_args(argv)
    if args.self_test:
        return self_test(args)
    return verify(args.cases, args.workbook, args.plan, args.export)


if __name__ == "__main__":
    sys.exit(main())
