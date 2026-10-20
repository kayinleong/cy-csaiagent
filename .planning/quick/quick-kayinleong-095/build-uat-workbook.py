"""Build the tester workbook from uat-cases.json.

The workbook has one sheet, `Test Cases`, holding only the sections named in
`meta.tester_sections` (Finder, new project data, Reply and Coach). The JSON stays the full
catalogue; the cross-cutting (X) cases and the exploratory slots are kept in it
for a later round but are not written here. Every cell is derived from the JSON.
The internal `source` field is never written. There are no formulas. Layout
mirrors docs/Phase-1-Test-Cases.xlsx.

Usage: python3 build-uat-workbook.py [--cases PATH] [--out PATH]
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]  # .planning/quick/quick-kayinleong-095 -> repo root
DEFAULT_CASES = HERE / "uat-cases.json"
DEFAULT_OUT = ROOT / "docs" / "UAT-Output-Test-Cases.xlsx"

FONT = "Arial"
INK = "1F2937"
INK_STRONG = "111827"
INK_BODY = "374151"
GREY = "6B7280"
BAND_FILL = PatternFill("solid", fgColor="D9F99D")
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
YELLOW = PatternFill("solid", fgColor="FFF9C4")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIORITY_COLOR = {"P0": "B91C1C", "P1": "92400E", "P2": INK_BODY}

LANGS = ["EN", "BM", "中文"]
ZIP_STAMP = (2026, 10, 19, 0, 0, 0)  # fixed entry time, so a rebuild is byte-identical


class FixedTimeZip(ZipFile):
    """ZipFile whose entries all carry ZIP_STAMP instead of the time of the build."""

    def writestr(self, zinfo_or_arcname, data, compress_type=None, compresslevel=None):
        if not isinstance(zinfo_or_arcname, ZipInfo):
            info = ZipInfo(filename=zinfo_or_arcname, date_time=ZIP_STAMP)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zinfo_or_arcname = info
        else:
            zinfo_or_arcname.date_time = ZIP_STAMP
        super().writestr(zinfo_or_arcname, data, compress_type, compresslevel)

    def write(self, filename, arcname=None, compress_type=None, compresslevel=None):
        # openpyxl writes the worksheet to a temp file and adds it with write(), which would stamp the file's mtime
        self.writestr(arcname or str(filename), Path(filename).read_bytes())

TC_HEADERS = ["ID", "Area", "Test Case", "Priority", "Preconditions", "Steps", "Expected Result", "Result",
              "Actual / Comments", "Known Status", "Lang", "Type", "Runs", "Run by", "Double-rate"]
TC_WIDTHS = [7, 16, 26, 9, 24, 40, 44, 12, 30, 34, 8, 13, 7, 15, 11]
RESULT_LIST = '"Pass,Fail,Blocked,Not Run"'


# ---------------------------------------------------------------- helpers
def font(size=10, bold=False, color=INK_BODY):
    return Font(name=FONT, size=size, bold=bold, color=color)


def wrap_top(horizontal=None):
    return Alignment(vertical="top", wrap_text=True, horizontal=horizontal)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def header_cells(ws, row, values, size=10, center=True):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = font(size, True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center" if center else "top", wrap_text=True)
        cell.border = BORDER


def band_row(ws, row, text, last_col, height=22):
    ws.cell(row=row, column=1, value=text)
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = BAND_FILL
        cell.font = font(11, True, INK)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.row_dimensions[row].height = height


def compose_expected(case):
    e = case["expected"]
    return "MUST: {}\nMUST NOT: {}\nALSO ACCEPTABLE: {}".format(e["must"], e["must_not"], e["also_acceptable"])


def compose_known(case, phrases):
    return "\n".join(phrases[k] for k in case["known_status_keys"])


def tester_sections(data):
    """Sections written to the workbook, in the order the JSON lists them."""
    wanted = data["meta"]["tester_sections"]
    known = [s["letter"] for s in data["sections"]]
    missing = [w for w in wanted if w not in known]
    if missing:
        raise SystemExit("meta.tester_sections names unknown sections: {}".format(missing))
    return [s for s in data["sections"] if s["letter"] in wanted]


# ---------------------------------------------------------------- Test Cases
def build_test_cases(wb, data):
    ws = wb.active
    ws.title = "Test Cases"
    ws.sheet_view.showGridLines = False
    set_widths(ws, TC_WIDTHS)
    header_cells(ws, 1, TC_HEADERS)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    phrases = data["known_status_phrases"]

    layout = {"case_rows": [], "bands": []}
    dv = DataValidation(type="list", formula1=RESULT_LIST, allow_blank=True)
    ws.add_data_validation(dv)

    def style_case_row(r):
        for c in range(1, 16):
            cell = ws.cell(row=r, column=c)
            cell.font = font(9, c == 1, INK_STRONG if c == 1 else INK_BODY)
            cell.alignment = wrap_top("center" if c in (8, 11, 13, 15) else None)
            cell.border = BORDER
            if c in (8, 9):
                cell.fill = YELLOW
        ws.cell(row=r, column=10).font = font(9, False, GREY)

    r = 1
    for section in tester_sections(data):
        r += 1
        band_row(ws, r, section["band"], 15)
        layout["bands"].append(r)
        sec_cases = [c for c in data["cases"] if c["section"] == section["letter"]]
        for lang in LANGS:
            for case in [c for c in sec_cases if c["lang"] == lang]:
                r += 1
                values = [case["id"], case["area"], case["title"], case["priority"], case["preconditions"],
                          case["steps"], compose_expected(case), "Not Run", None, compose_known(case, phrases),
                          case["lang"], case["type"], case["runs"], case["run_by"],
                          "Yes" if case["double_rate"] else None]
                for c, v in enumerate(values, start=1):
                    ws.cell(row=r, column=c, value=v)
                style_case_row(r)
                ws.cell(row=r, column=4).font = font(9, False, PRIORITY_COLOR[case["priority"]])
                dv.add("H{}".format(r))
                layout["case_rows"].append(r)
    layout["last_row"] = r
    return layout


# ---------------------------------------------------------------- main
def build(cases_path, out_path):
    data = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    wb = Workbook()
    tc = build_test_cases(wb, data)
    stamp = datetime(2026, 10, 19, 0, 0, 0)  # fixed, so a rebuild changes nothing
    wb.properties.creator = "UAT build"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    # wb.save() stamps `modified` with the current time, which would change the file on every rebuild.
    # ExcelWriter writes the fixed stamp above instead, so the same JSON gives the same workbook.
    with FixedTimeZip(str(out_path), "w", ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(wb, archive).save()
    return data, tc


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", default=str(DEFAULT_CASES))
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    args = ap.parse_args(argv)
    data, tc = build(args.cases, args.out)
    print("wrote {}".format(args.out))
    print("Test Cases: {} cases in sections {}, {} band rows, last row {}".format(
        len(tc["case_rows"]), "".join(data["meta"]["tester_sections"]), len(tc["bands"]), tc["last_row"]))


if __name__ == "__main__":
    sys.exit(main())
