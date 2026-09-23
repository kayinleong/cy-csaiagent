#!/usr/bin/env python3
"""
quick-kayinleong-092 — fold the Drive document text into the quick-091 workbook.

Still ONE tab, `Raw Data`. Each Drive document gets its own pair of columns:

    Drive Doc 01 Name | Drive Doc 01 Text | Drive Doc 02 Name | Drive Doc 02 Text | ...

Documents are assigned to slots in filename order, so slot N means the same
position for every project. Excel caps a cell at 32,767 characters, so a document
longer than that continues into `... Text (cont. 2)`, `(cont. 3)` and so on —
only as many continuation columns as the widest document in that slot needs.
Nothing is truncated.

The KB text comes from `export-kb-text.ts` (already-ingested Drive docs read back
out of Firestore). WhatsApp-bucket docs are excluded: they are group-chat
transcripts, not Drive collateral.

    uv run --with openpyxl python attach-kb-text.py \
        --xlsx D2-Property-Raw-Data.xlsx --kb <scratch>/kb-text.json --out <out>.xlsx
"""

import argparse
import json
import re
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CELL_CAP = 32000  # Excel's hard limit is 32,767; leave headroom.

# XML 1.0 forbids these control characters, so no .xlsx can hold them. Some of the
# OCR'd / badly-encoded PDFs carry them. Stripping them is not a content decision --
# the alternative is that the file cannot be written at all.
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(text):
    return ILLEGAL.sub("", text) if isinstance(text, str) else text


def split_cap(text, cap=CELL_CAP):
    text = clean(text)
    return [text[i:i + cap] for i in range(0, len(text), cap)] or [""]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="the quick-091 workbook")
    ap.add_argument("--kb", required=True, help="kb-text.json from export-kb-text.ts")
    ap.add_argument("--out", required=True)
    ap.add_argument("--include-whatsapp", action="store_true",
                    help="also fold in the WhatsApp transcript bucket (23.5 MB)")
    args = ap.parse_args()

    src = load_workbook(args.xlsx)["Raw Data"]
    headers = [c.value for c in src[1]]
    rows = [dict(zip(headers, [c.value for c in r])) for r in src.iter_rows(min_row=2)]
    names = {r["Project"] for r in rows}

    kb = [d for d in json.load(open(args.kb))["docs"] if d["chars"] > 0]

    # Two joins. Normal Drive docs carry the project in `project` (to-kb.ts writes
    # the title as "<project> — <filename>"). The WhatsApp bucket inverts that --
    # project is literally "WhatsApp" and the project name sits in `fileName`.
    by_project = defaultdict(list)
    for d in kb:
        if d["project"] == "WhatsApp":
            if args.include_whatsapp and d["fileName"] in names:
                by_project[d["fileName"]].append(d)
        elif d["project"] in names:
            by_project[d["project"]].append(d)
    for v in by_project.values():
        v.sort(key=lambda x: (x["fileName"] or "").lower())

    max_docs = max((len(v) for v in by_project.values()), default=0)

    # How many text columns each slot needs = widest document in that slot.
    slot_parts = []
    for i in range(max_docs):
        widest = max((len(v[i]["text"]) for v in by_project.values() if len(v) > i), default=0)
        slot_parts.append(max(1, -(-widest // CELL_CAP)))

    out_headers = list(headers) + ["Drive Doc Count", "Drive Doc Titles"]
    for i, parts in enumerate(slot_parts, start=1):
        out_headers.append(f"Drive Doc {i:02d} Name")
        out_headers.append(f"Drive Doc {i:02d} Text")
        for p in range(2, parts + 1):
            out_headers.append(f"Drive Doc {i:02d} Text (cont. {p})")

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Raw Data"
    ws.append(out_headers)
    head_fill = PatternFill("solid", fgColor="1F3864")
    kb_fill = PatternFill("solid", fgColor="4472C4")
    for i, h in enumerate(out_headers, start=1):
        c = ws.cell(row=1, column=i)
        c.fill = kb_fill if str(h).startswith("Drive Doc") else head_fill
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    embedded = 0
    for r in rows:
        docs = by_project.get(r["Project"], [])
        line = [clean(r.get(h, "")) for h in headers]
        line.append(len(docs))
        line.append(clean("\n".join(d["fileName"] + (" (OCR)" if d["ocr"] else "") for d in docs)))
        for i, parts in enumerate(slot_parts):
            d = docs[i] if i < len(docs) else None
            line.append(clean((d["fileName"] + (" (OCR)" if d["ocr"] else ""))) if d else "")
            chunks = split_cap(d["text"]) if d else [""]
            if d:
                embedded += len(d["text"])
            for p in range(parts):
                line.append(chunks[p] if p < len(chunks) else "")
        ws.append(line)

    # Only size the original columns; the text columns stay narrow so the sheet
    # stays navigable, and wrap is left off so Excel does not try to lay out
    # 32,000-character cells.
    for i, h in enumerate(out_headers, start=1):
        letter = get_column_letter(i)
        if str(h).startswith("Drive Doc"):
            ws.column_dimensions[letter].width = 18 if "Name" in str(h) else 40
        else:
            longest = max([len(str(h))] + [len(str(r.get(h, "")).split("\n")[0]) for r in rows])
            ws.column_dimensions[letter].width = min(max(12, longest + 2), 52)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(out_headers))}{len(rows) + 1}"

    wb.save(args.out)
    covered = sum(1 for r in rows if by_project.get(r["Project"]))
    print(f"[attach] {len(rows)} rows x {len(out_headers)} columns -> {args.out}")
    print(f"[attach] {covered}/{len(rows)} projects carry Drive docs; "
          f"max {max_docs} docs on one project")
    print(f"[attach] {embedded:,} characters of document text embedded")


if __name__ == "__main__":
    main()
