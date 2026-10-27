"""Build the run-3 workbook. usage: python3 -E -B build3.py <run3 dir> <out.xlsx>"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

d = Path(sys.argv[1])
out = sys.argv[2]
scored = json.loads((d / "scored.json").read_text())
results = json.loads((d / "results.json").read_text())
run2 = json.loads((d / "run2.json").read_text())
search = json.loads((d / "search3.json").read_text())
overrides = json.loads((d / "overrides.json").read_text()) if (d / "overrides.json").exists() else {}

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
B = Font(name=F, bold=True)
N = Font(name=F)
HFILL = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")
FILLS = {
    "PASS": "C6EFCE", "Pass": "C6EFCE", "PARTIAL": "FFEB9C", "Partial": "FFEB9C", "NOT STATED": "FFEB9C",
    "DIFFERS": "FFC7CE", "Check": "FFC7CE", "Fail": "FFC7CE", "HELD": "D9D9D9", "Needs ruling": "DDEBF7",
}
COLS = ["Size", "Rooms", "Nett Price", "Gross Price", "Area", "Developer", "Land Title", "Maintenance Fee", "Tenure",
        "Total Units", "Total Floors", "Completion / VP", "Top Reasons Why", "Furnishing", "Car Park Allocation"]


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.alignment = H, HFILL, Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def colour(ws, rng):
    for k, hexv in FILLS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % k], fill=PatternFill("solid", fgColor=hexv)))


# apply manual overrides (cell-level, after reading the answers)
for r in scored:
    for col, ov in overrides.get(r["id"], {}).items():
        if col == "_note":
            r["notes"].append(ov)
            continue
        r["cells"][col] = ov["verdict"]
        r["details"][col] = ov["why"]
    vals = list(r["cells"].values())
    if r.get("overall") not in ("ERROR", "NO SHEET ROW"):
        r["overall"] = "Check" if "DIFFERS" in vals else "Partial" if any(v in ("NOT STATED", "PARTIAL") for v in vals) else "Pass"

wb = openpyxl.Workbook()
s = wb.active
s.title = "Summary"
fc = wb.create_sheet("Field Check (run 3)")
cmp_ = wb.create_sheet("Run 2 vs Run 3")
sc = wb.create_sheet("Search Cases")
fnd = wb.create_sheet("Findings")
ans = wb.create_sheet("Answers")

# ── Field Check
fc["A1"] = "Client Sheet Field Check, RUN 3 (27 Oct 2026): every property's answer scored against the client Excel (_D2 Project Info.xlsx)"
fc["A1"].font = Font(name=F, bold=True, size=12)
fc["A2"] = ("PASS = matches the Excel · PARTIAL = some layouts match, some missing · NOT STATED = the Excel has a value, the answer "
            "does not give it · DIFFERS = the answer gives a different value · HELD = held back from the import pending the client's "
            "confirmation · — = the Excel has no value")
fc["A2"].font = Font(name=F, italic=True)
hdr = ["#", "Project (Excel name)", "Overall"] + COLS + ["Fact sheet attached", "Run 2 overall (client table)", "Notes"]
header(fc, 4, hdr, [5, 36, 10] + [11] * len(COLS) + [11, 14, 90])
for i, r in enumerate(scored, start=5):
    row = [int(r["id"][1:]), r["name"], r["overall"]] + [r["cells"].get(c, "") for c in COLS]
    row += ["Yes" if r.get("factSheetAttached") else "No", (run2["field"].get(r["id"]) or {}).get("Overall", "")]
    notes = ["%s: %s" % (k, v) for k, v in r["details"].items()] + r["notes"]
    row.append(" | ".join(notes))
    for j, v in enumerate(row, start=1):
        c = fc.cell(row=i, column=j, value=v)
        c.font, c.alignment = N, WRAP
last = 4 + len(scored)
colour(fc, "C5:%s%d" % (get_column_letter(3 + len(COLS) + 2), last))
fc.freeze_panes = "D5"

# ── Run 2 vs Run 3
header(cmp_, 1, ["#", "Project", "Run 2 (client test table)", "Run 3 (Excel)", "Same?"], [5, 40, 22, 14, 8])
for i, r in enumerate(scored, start=2):
    r2 = (run2["field"].get(r["id"]) or {}).get("Overall", "")
    r2n = "Pass" if str(r2).startswith("Pass") else r2
    for j, v in enumerate([int(r["id"][1:]), r["name"], r2, r["overall"], "Yes" if r2n == r["overall"] else "No"], start=1):
        cmp_.cell(row=i, column=j, value=v).font = N
colour(cmp_, "C2:D%d" % (1 + len(scored)))

# ── Search cases
header(sc, 1, ["Case", "Prompt (sent exactly)", "Run 3", "Run 2", "What the Finder returned", "Verdict notes"], [7, 50, 12, 12, 70, 70])
for i, (cid, v) in enumerate(search.items(), start=2):
    row = [cid.replace("S-", ""), results[cid]["prompt"], v["result"], run2["search"].get(cid, {}).get("result", ""), v["returned"], v["notes"]]
    for j, x in enumerate(row, start=1):
        c = sc.cell(row=i, column=j, value=x)
        c.font, c.alignment = N, WRAP
colour(sc, "C2:D%d" % (1 + len(search)))

# ── Findings
header(fnd, 1, ["#", "Severity", "Finding", "Evidence", "Recommended action"], [4, 10, 55, 70, 60])
FINDINGS = [
    ("Medium", "Eaton Residences: the answer leads with 632 units instead of the Excel's 960",
     "The client fact sheet (960 units: Tower A 492 + Tower B 468) was in the bot's context. An older sales kit says '52 floors, Total Units: 632units'; the bot stated 632 and gave 960 only as a side note, against the new 'fact sheet is the authority' rule.",
     "Ask the client which is right (632 may be one tower). If the Excel stands and this recurs, put the fact-sheet values into the project record itself rather than relying on the prompt rule."),
    ("Medium", "Yanu Hills: the Excel's Area disagrees with its own Land Title",
     "Excel Area: 'Taman Melawati, Kuala Lumpur.' Excel Land Title: 'Selangor Individual'. The brochure says Taman Nadayu, Melawati, Selangor, and that is what the bot said.",
     "Ask the client to confirm the area (Kuala Lumpur or Selangor)."),
    ("Low", "The Lantern Bangsar: layout labels differ from the Excel",
     "Excel: 'Type A: 570 sf'. The bot: 'Type A1: 570 sq ft' plus a 'Type A2: 561 sq ft' that is not in the Excel. Claim 096 already flagged Lantern's near-duplicate Type A / A1 layouts.",
     "Ask the client whether Type A, A1 and A2 are separate layouts."),
    ("High (client data)", "The Excel's Developer column is still shifted in rows 91-93 and 114-125",
     "Each of those cells holds a neighbouring project's developer (e.g. St. Regis shows 'UOA Development Bhd'). Answers now state the Excel's value, so these rows PASS against the Excel while being wrong in the real world.",
     "Ask the client to correct those Developer cells; re-import with the same pipeline afterwards."),
    ("Info", "Search case F11 changed because of the new Binastra Cochrane",
     "Binastra Cochrane is in Cheras with no price in the Excel; the agreed search rule (quick-kayinleong-085) admits unpriced projects to a budget search, so the 3-bedroom/RM650k Cheras lead now gets it, with a note to confirm price and bedrooms.",
     "Decide whether unpriced projects should appear for budgeted leads, or ask the client for Binastra Cochrane's prices."),
    ("Info", "Values held back from the import (HELD)",
     "23 values in HELD-BACK.md (claim 096) were not loaded pending the client's confirmation; 8 cells in this test touch them and are not scored.",
     "Send HELD-BACK.md to the client."),
    ("Info", "Exsim Project has no facts in the Excel",
     "Its row has no usable values, so there is no fact sheet (the only project without one) and nothing to score.",
     "None unless the client adds details."),
]
for i, f in enumerate(FINDINGS, start=2):
    for j, v in enumerate((i - 1,) + f, start=1):
        c = fnd.cell(row=i, column=j, value=v)
        c.font, c.alignment = N, WRAP

# ── Answers
header(ans, 1, ["ID", "Project / case", "Seconds", "Finish", "Answer (as returned)"], [7, 36, 9, 9, 140])
order = list(search.keys()) + [r["id"] for r in scored]
for i, k in enumerate(order, start=2):
    r = results[k]
    t = r.get("text") or r.get("error", "")
    for j, x in enumerate([k, r.get("name") or k, round(r.get("ms", 0) / 1000, 1), r.get("finishReason", "error"), t[:32000]], start=1):
        c = ans.cell(row=i, column=j, value=x)
        c.font, c.alignment = N, WRAP

# ── Summary
rows = [
    ("D2 Finder: client Excel field check, run 3 (after quick-kayinleong-101)", None),
    (None, None),
    ("Date", "27 Oct 2026"),
    ("Oracle", "_D2 Project Info.xlsx (sha256 8b7c1c9a…), as normalised by the import (81 mapped projects + the 3 added in this claim)"),
    ("How it was run", "Script-driven: the real Finder agent (same system prompt, tools, 12-step budget, 32k output, PDPA scrubbing as /api/chat), "
                       "no browser and no sign-in. Model gemini-3.5-flash via the dev-only LOCAL_MODEL_OVERRIDE, as runs 1-2."),
    ("Code under test", "branch quick-kayinleong-101: projectDetail always attaches the project's own client fact sheet; the sheet is the authority"),
    ("Data under test", "live Firestore after the 3 new projects (Binastra Cochrane, Lagenda Tunku, Ashbourne) and their fact sheets were added"),
    ("Prompt per property", results[order[len(search)]]["prompt"].replace(order and results[order[len(search)]]["name"] or "", "<project>")),
    (None, None),
    ("Properties tested", "=COUNTA('Field Check (run 3)'!B5:B%d)" % last),
    ("Overall Pass", "=COUNTIF('Field Check (run 3)'!C5:C%d,\"Pass\")" % last),
    ("Overall Partial", "=COUNTIF('Field Check (run 3)'!C5:C%d,\"Partial\")" % last),
    ("Overall Check (a value differs)", "=COUNTIF('Field Check (run 3)'!C5:C%d,\"Check\")" % last),
    ("Field cells PASS", "=COUNTIF('Field Check (run 3)'!D5:R%d,\"PASS\")" % last),
    ("Field cells PARTIAL", "=COUNTIF('Field Check (run 3)'!D5:R%d,\"PARTIAL\")" % last),
    ("Field cells NOT STATED", "=COUNTIF('Field Check (run 3)'!D5:R%d,\"NOT STATED\")" % last),
    ("Field cells DIFFERS", "=COUNTIF('Field Check (run 3)'!D5:R%d,\"DIFFERS\")" % last),
    ("Field cells HELD", "=COUNTIF('Field Check (run 3)'!D5:R%d,\"HELD\")" % last),
    ("Fact sheet attached", "=COUNTIF('Field Check (run 3)'!S5:S%d,\"Yes\")" % last),
    ("Search cases Pass", "=COUNTIF('Search Cases'!C2:C%d,\"Pass\")" % (1 + len(search))),
    (None, None),
    ("Still in the Excel", "The Developer column is shifted in rows 91-93 and 114-125 (each cell holds a neighbouring project's developer). "
                           "Answers now state the Excel's value, so those rows PASS here while still being wrong in the real world."),
    ("Held back", "23 values the import held back pending the client's confirmation (HELD-BACK.md) are shown as HELD, not scored."),
]
for i, (k, v) in enumerate(rows, start=1):
    a = s.cell(row=i, column=1, value=k)
    b = s.cell(row=i, column=2, value=v)
    a.font = Font(name=F, bold=True, size=13) if i == 1 else B
    b.font, b.alignment = N, WRAP
s.column_dimensions["A"].width = 34
s.column_dimensions["B"].width = 120

wb.calculation.fullCalcOnLoad = True
wb.save(out)
print("saved", out)
