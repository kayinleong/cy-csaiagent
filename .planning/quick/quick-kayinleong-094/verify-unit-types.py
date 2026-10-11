"""Check the transcription in unit-types.json and the workbook built from it.

The transcription is checked against facts that are independent of each other
on the page: each sqft figure against its printed sqm, the layout's bedroom
count against the bedrooms drawn on the plan, mirrored types across the two
towers against each other, and the range against the brochure's own
"18 layouts ranging from 904 sf to 2,982 sf".
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SQFT_PER_SQM = 10.7639
MIRRORS = [("S1", "S2"), ("A1", "A5"), ("A2", "A6"), ("B1", "B5"), ("B2", "B6"), ("B3", "B7")]
LISTING_RE = re.compile(r"^Type [A-Z]\d \d+ sqft \d+sqm \dbr(\+1){0,2}$")

failures = []


def check(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    if not ok:
        failures.append(msg)


data = json.loads((HERE / "unit-types.json").read_text())
units = data["units"]
by_type = {u["type"]: u for u in units}

check(len(units) == 18, f"18 unit types (got {len(units)})")
check(len(by_type) == 18, "type codes are unique")
for t in ("A", "B"):
    nos = sorted(u["key_plan_no"] for u in units if u["tower"] == t)
    check(nos == sorted(["1", "2", "3", "3a", "5", "6", "7", "8", "9"]), f"Tower {t} key-plan numbers are 1-3, 3a, 5-9 once each")

sizes = [u["sqft"] for u in units]
check(min(sizes) == 904 and max(sizes) == 2982, "range is 904-2,982 sqft as stated on brochure p.19")

for u in units:
    sqm_from_sqft = u["sqft"] / SQFT_PER_SQM
    check(abs(sqm_from_sqft - u["sqm"]) < 1, f"{u['type']}: {u['sqft']} sqft agrees with printed {u['sqm']} sqm ({sqm_from_sqft:.1f})")
    br = int(re.match(r"(\d)BR", u["layout"]).group(1))
    drawn = sum("Bedroom" in r for r in u["rooms"])
    check(br == drawn, f"{u['type']}: layout {u['layout']} has {br} BR, plan draws {drawn} bedrooms")
    check(len(u["rooms"]) == len(set(u["rooms"])), f"{u['type']}: no duplicate room labels")

for a, b in MIRRORS:
    ua, ub = by_type[a], by_type[b]
    same = all(ua[k] == ub[k] for k in ("sqft", "sqm", "layout", "key_plan_no")) and sorted(ua["rooms"]) == sorted(ub["rooms"])
    check(same, f"mirror pair {a}/{b} matches on size, layout, key-plan no. and rooms")

# --- workbook readback -------------------------------------------------
wb = load_workbook(HERE / "DC-Residensi-Unit-Types.xlsx")
check(wb.sheetnames == ["Unit Types", "Paste Into Raw Data", "Notes"], f"sheets are {wb.sheetnames}")
ws = wb["Unit Types"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
check(len(rows) == 18, f"Unit Types has 18 data rows (got {len(rows)})")
listings = [r[0] for r in rows]
check(all(LISTING_RE.match(s) for s in listings), "every listing matches the client format")
check("Type B2 1851 sqft 172sqm 2br+1+1" in listings, "client's own example appears verbatim")
for r, u in zip(rows, units):
    expected = f"Type {u['type']} {u['sqft']} sqft {u['sqm']}sqm {u['layout'].lower()}"
    check(r[0] == expected and r[3] == u["type"] and r[4] == u["sqft"] and r[5] == u["sqm"] and r[6] == u["layout"],
          f"row {u['type']} matches the transcription")
    cells_rooms = sorted(x.strip() for col in r[7:11] if col and col != "-" for x in col.split(","))
    check(cells_rooms == sorted(u["rooms"]), f"row {u['type']} room columns hold every plan label exactly once")

paste = wb["Paste Into Raw Data"]["A2"].value.split("\n")
check(paste == listings, "paste sheet lines equal the 18 listings, in order")

fonts = {c.font.name for s in wb for row in s.iter_rows() for c in row if c.value is not None}
check(fonts == {"Arial"}, f"fonts used: {fonts}")
check(not any(isinstance(c.value, str) and c.value.startswith("=") for s in wb for row in s.iter_rows() for c in row),
      "no formulas (nothing needs recalculation)")

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
