"""Build the DC Residensi unit-type workbook from unit-types.json.

The brochure is image-only, so unit-types.json is a hand transcription of the
floor-plan pages. This script only formats it: the client-format listing is a
string template over the transcribed fields, and the room columns are a
label-for-label split of each plan's room list. Nothing is calculated.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
OUT = HERE / "DC-Residensi-Unit-Types.xlsx"

# Every label on the plans lands in exactly one bucket; an unknown label fails
# the build instead of being dropped silently.
EXTRA_ROOMS = {"Study Room", "Utility", "Maids Room", "Maid's Room", "Family Room", "Walk In Wardrobe"}
BATHS = {"Master Bath", "Bath 2", "Bath 3", "Pwdr Room", "WC"}
OTHER = {"Living", "Dining", "Dry Kitchen", "Wet Kitchen", "Foyer", "Lobby", "Private Lobby", "Yard", "Balcony"}

BROCHURE_PAGES = {14: "pp.24-25 (Tower A floor plans)", 15: "pp.26-27 (Tower B floor plans)"}

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def listing(u):
    # Client format, e.g. "Type B2 1851 sqft 172sqm 2br+1+1"
    return f"Type {u['type']} {u['sqft']} sqft {u['sqm']}sqm {u['layout'].lower()}"


def split_rooms(rooms):
    beds, extra, baths, other = [], [], [], []
    for r in rooms:
        if "Bedroom" in r:
            beds.append(r)
        elif r in EXTRA_ROOMS:
            extra.append(r)
        elif r in BATHS:
            baths.append(r)
        elif r in OTHER:
            other.append(r)
        else:
            sys.exit(f"unclassified room label: {r!r}")
    return beds, extra, baths, other


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32


def main():
    data = json.loads((HERE / "unit-types.json").read_text())
    units = data["units"]

    wb = Workbook()

    # --- Sheet 1: one row per unit type ---------------------------------
    ws = wb.active
    ws.title = "Unit Types"
    headers = [
        "Listing (client format)", "Project", "Tower", "Type", "Built-Up (sqft)", "Built-Up (sqm)",
        "Layout (as per brochure)", "Bedrooms (as labelled)", "Study / Utility / Maid's / Family (as labelled)",
        "Bathrooms & WC (as labelled)", "Other spaces (as labelled)", "Unit No. on Key Plan", "Source",
    ]
    widths = [36, 22, 8, 7, 11, 11, 13, 30, 34, 36, 52, 12, 34]
    ws.append(headers)
    for u in units:
        beds, extra, baths, other = split_rooms(u["rooms"])
        ws.append([
            listing(u), data["project"], f"Tower {u['tower']}", u["type"], u["sqft"], u["sqm"], u["layout"],
            ", ".join(beds), ", ".join(extra) or "-", ", ".join(baths), ", ".join(other), u["key_plan_no"],
            f"Brochure {BROCHURE_PAGES[u['pdf_page']]}",
        ])
    style_header(ws, len(headers))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.column == 1))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        row[4].number_format = "0"  # client format writes 1851, not 1,851
        row[5].number_format = "0"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # --- Sheet 2: all 18 lines in one cell, for the Raw Data column -----
    ws2 = wb.create_sheet("Paste Into Raw Data")
    ws2.append(["Built-Up sizes & Layout (one line per type)"])
    ws2.append(["\n".join(listing(u) for u in units)])
    style_header(ws2, 1)
    ws2.column_dimensions["A"].width = 44
    ws2["A2"].font = Font(name=FONT)
    ws2["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    ws2.row_dimensions[2].height = 15 * len(units)

    # --- Sheet 3: source and caveats ------------------------------------
    ws3 = wb.create_sheet("Notes")
    notes = [
        ("Source", data["source"]),
        ("How extracted", data["source_note"]),
        ("Completeness check", "18 types listed. Brochure p.19 states \"18 layouts ranging from 904 sf to 2,982 sf\" - "
                               "matches the smallest (S1/S2, 904 sqft) and largest (E2, 2,982 sqft) rows here."),
        ("Room columns", "Copied label-for-label from each floor plan. The brochure does not say which room each \"+1\" "
                         "refers to, so the Study / Utility / Maid's / Family column lists every such room drawn on the plan."),
        ("Unit No. on Key Plan", "The number printed under each type on the tower key plan (there is no unit 4; "
                                 "the brochure uses 3a)."),
        ("Not in this brochure", "Penthouses, car park bays per type, prices, facing/view per unit."),
        ("Differs from the current Skool write-up",
         "Skool says \"1+1 Bedrooms: 1,100 sqft\" - brochure 1BR+1 types are 1,152-1,194 sqft. "
         "Skool says \"2+1 Bedrooms: 1,600-1,800 sqft\" - brochure has 2BR+1 at 1,668 and 2BR+1+1 at 1,765-1,862. "
         "Skool says \"3+1 Bedrooms: 2,700-2,900 sqft\" - brochure has 3BR+1 (C1) at 2,077 and 3BR+1+1 at 2,702 and 2,982. "
         "Skool's \"Penthouses: 2,900-4,855 sqft\" does not appear in this brochure."),
    ]
    ws3.append(["Item", "Detail"])
    for n in notes:
        ws3.append(list(n))
    style_header(ws3, 2)
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 110
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        row[0].font = Font(name=FONT, bold=True)
        row[1].font = Font(name=FONT)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(OUT)
    print(f"wrote {OUT.name}: {len(units)} unit types")


if __name__ == "__main__":
    main()
